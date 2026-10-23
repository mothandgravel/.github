#!/usr/bin/env python3
"""Build the Finnish public account from pinned publisher downloads, without network.

Usage: python scripts/build_finland_data.py artifacts/research/finland-energy [--check]
Read-only XLSX extraction requires openpyxl==3.1.5. Original files are not modified.
"""
import argparse
import hashlib
import json
import math
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'docs/assets/data-center/finland.json'
MANIFEST = ROOT / 'docs/design/finland-provenance.json'
INPUTS = {
    'statfin-13rb.json': ('b01492f4309c65dcfa18421ff55d7d83bdb38810d26cfc890f0b2776847b8e38', 'https://pxdata.stat.fi/PxWeb/api/v1/en/StatFin/ehi/13rb.px'),
    'statfin-13rb-query.json': ('d9a669e5b7a0c0da3508aaa57a5ff4372224bbc68a98d9bab00b6c7fc365dadb', None),
    'grid.xlsx': ('b3dcec0cdf043bf218f5e78b995e93cefd27e21c8911af65015986089e8beb9e', 'https://energia.fi/wp-content/uploads/2026/10/tuntidata_2020-2025.xlsx'),
    'heat.xlsx': ('dcafa61b2844cd2aa16121604f5818b1697425565c11040def1d6443407c78a1', 'https://energia.fi/wp-content/uploads/2026/06/KL_Vuositaulukot_24_paivitetty_20260612.xlsx'),
    'nivos-2025.pdf': ('381a60a127f1fbabd7e6fa9592876e3e7592d05b3e8e6ebd804d4245b1000d84', 'https://nivos.fi/wp-content/uploads/2026/04/Nivos-Vesi-ja-Lampo-Oy-tilinpaatos-ja-toimintakertomus-2025.pdf'),
    'nivos-origin.html': ('1a05945730578ea32ac8e62f9cd1762388ead5f59fe311de619c96a6ef9555a2', 'https://nivos.fi/kaukolampo/kaukolammon-alkupera/'),
}
METHODS = {
    'bill': 'Published monthly consumer-band unit prices. Energy and distribution exclude tax; the tax channel includes electricity excise and VAT. Published total retained. Annualised illustration = authored annual kWh × selected-month c/kWh / 100. No extra VAT, invented load profile or causal price model.',
    'grid': 'Energiateollisuus 2024 worksheet. Hour labels 1–24 interpreted as hour-ending slots in publisher fixed UTC+2 normal time; end = source date + hour, start = end − 1 h. No daylight-saving adjustment. MW are treated as hourly means. Energy = sum(MW × 1 h); no trapezoidal integration. Exports are signed negative in the source. Balance = production + import + signed export − consumption. These are national flows, not available capacity.',
    'heat': 'Published 2024 Nivos company/municipal row 114. Thermal GWh retained separately from electrical MWh. Recovered heat / heat-pump output is a combined source category, not pure server waste heat. Residential use is reported directly, not allocated by fuel mix. 2025 company sales (43.1 GWh) transcribed from report PDF page 3; central-network origin share (41%) transcribed from the pinned origin disclosure. Different scopes; no multiplication across them.',
    'branch': 'Authored extra whole-facility MW is added arithmetically to a historical national demand observation. The unmet-by-fixed-historical-supply increment is a requirement to investigate, not a forecast of imports, generation, blackouts, grid capacity or price. No PUE multiplier.',
}


def emit(path, data, check):
    text = json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n'
    if check:
        if not path.exists() or path.read_text(encoding='utf-8') != text:
            raise SystemExit(f'Generated output differs: {path}')
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding='utf-8', newline='\n')
    return hashlib.sha256(text.encode()).hexdigest()


def build(folder):
    import openpyxl
    for name, (sha, _) in INPUTS.items():
        assert hashlib.sha256((folder/name).read_bytes()).hexdigest() == sha, name
    raw = json.loads((folder/'statfin-13rb.json').read_bytes())
    assert raw['id'] == ['timeperiod_m', 'energia_48_20220905', 'energia_49_20220905', 'contentscode']
    assert raw['size'] == [30, 4, 5, 1]
    months = list(raw['dimension']['timeperiod_m']['category']['index'])
    bands = raw['dimension']['energia_49_20220905']['category']['label']
    prices = []
    max_price_residual = 0
    for t, month in enumerate(months):
        for b, band in enumerate(bands):
            indices = [(t*4+c)*5+b for c in range(4)]
            values = [raw['value'][i] for i in indices]
            assert all(isinstance(v, (int, float)) and math.isfinite(v) and v >= 0 for v in values)
            max_price_residual = max(max_price_residual, abs(sum(values[:3])-values[3]))
            prices.append({'month': month, 'band': band, 'energy': values[0], 'network': values[1], 'tax': values[2], 'total': values[3], 'valueIndices': indices})
    assert max_price_residual < .002
    workbook = openpyxl.load_workbook(folder/'grid.xlsx', read_only=True, data_only=True)
    sheet = workbook['2024']
    rows = list(sheet.values)
    assert 'Production' in rows[0][6] and 'Consumption' in rows[0][17]
    all_grid = []
    for r, row in enumerate(rows[1:], 2):
        y, m, d, h = row[:4]
        assert y == 2024 and 1 <= h <= 24
        end = datetime(y, m, d, tzinfo=timezone(timedelta(hours=2))) + timedelta(hours=h)
        vals = [row[c] for c in [17, 6, 15, 16, 8]]
        assert all(isinstance(v, (int, float)) and math.isfinite(v) for v in vals)
        assert vals[0] >= 0 and vals[1] >= 0 and vals[2] >= 0 and vals[3] <= 0 and vals[4] >= 0
        all_grid.append({'start': (end-timedelta(hours=1)).isoformat(), 'sourceDay': f'{y:04}-{m:02}-{d:02}', 'sourceHour': h, 'row': r, 'consumption': vals[0], 'production': vals[1], 'import': vals[2], 'exportSigned': vals[3], 'wind': vals[4]})
    assert len(all_grid) == 8784
    assert len({r['start'] for r in all_grid}) == 8784
    for a, b in zip(all_grid, all_grid[1:]):
        assert datetime.fromisoformat(b['start'])-datetime.fromisoformat(a['start']) == timedelta(hours=1)
    max_grid_residual = max(abs(r['production']+r['import']+r['exportSigned']-r['consumption']) for r in all_grid)
    assert max_grid_residual < .000001
    windows = []
    for key, name, start, end in [('winter', '1–7 January 2024', '2024-01-01', '2024-01-08'), ('summer', '1–7 July 2024', '2024-07-01', '2024-07-08')]:
        selected = [r for r in all_grid if start <= r['sourceDay'] < end]
        assert len(selected) == 168
        windows.append({'id': key, 'label': name, 'rows': selected, 'consumptionMWhDecimal': str(sum(Decimal(str(r['consumption'])) for r in selected))})
    heatbook = openpyxl.load_workbook(folder/'heat.xlsx', read_only=True, data_only=True)
    cells = []
    mappings = {
        'Taul1': {'fuelHeat': 'F148', 'recoveryAndHeatPump': 'G148', 'supply': 'I148', 'customerUse': 'J148', 'lossAndMeterDifference': 'L148', 'productionElectricityMWh': 'O148', 'distributionPumpElectricityMWh': 'P148'},
        'Taul5': {'residentialCustomers': 'F138', 'residentialUse': 'R138', 'industrialUse': 'S138', 'otherUse': 'T138', 'allCustomerUse': 'U138'},
    }
    for sn, mapping in mappings.items():
        s = heatbook[sn]
        assert 'Nivos' in s['E148' if sn == 'Taul1' else 'E138'].value
        for key, addr in mapping.items():
            val = s[addr].value
            source_label = s.cell(4, s[addr].column).value
            cells.append({'id': key, 'sheet': sn, 'cell': addr, 'sourceLabel': source_label, 'raw': val, 'value': val if isinstance(val, (int, float)) else None, 'unit': 'MWh' if 'Electricity' in key else 'customers' if key == 'residentialCustomers' else 'GWh'})
    h = {c['id']: c['value'] for c in cells}
    assert abs(h['fuelHeat']+h['recoveryAndHeatPump']-h['supply']) < 1e-8
    assert abs(h['supply']-h['customerUse']-h['lossAndMeterDifference']) < 1e-8
    assert abs(h['residentialUse']+h['industrialUse']+h['otherUse']-h['customerUse']) < 1e-8
    return {
        'schema': 'moth-gravel/finland-account@1', 'retrieved': '2026-10-23', 'methods': METHODS,
        'prices': {'source': 'Statistics Finland, energy prices / 13rb', 'updated': raw['updated'], 'unit': 'c/kWh', 'months': months, 'bandLabels': bands, 'labelNote': 'The source English band-A label says kW. Display labels use kWh for annual energy, consistently with the other annual-consumption bands. Original labels are retained here.', 'rows': prices, 'query': json.loads((folder/'statfin-13rb-query.json').read_bytes()), 'notes': raw['note'], 'maxComponentResidualCKWh': max_price_residual},
        'grid': {'source': 'Energiateollisuus ry / Sähkön tuntidata 2020–2025', 'sheet': '2024', 'unit': 'MW', 'sourceColumns': {'consumption': 'R', 'production': 'G', 'import': 'P', 'exportSigned': 'Q', 'wind': 'I'}, 'windows': windows, 'year': {'year': 2024, 'hours': 8784, 'consumptionMWhDecimal': str(sum(Decimal(str(r['consumption'])) for r in all_grid)), 'maxBalanceResidualMW': max_grid_residual}},
        'heat': {'source': 'Energiateollisuus ry, district heat statistics 2024; Nivos reporting', 'company': 'Nivos Vesi ja Lämpö Oy, Mäntsälä', 'scope2024': 'Company/municipal row 114; multiple networks; residential, industrial and other customers.', 'cells': cells, 'report2025': {'salesGWh': 43.1, 'scope': 'All company district-heat sales; residential breakdown unavailable in this extract.', 'pdfPage': 3, 'printedPage': 2, 'centralNetworkRecoveryPercent': 41, 'originScope': 'Mäntsälä central network only. Rounded origin share; do not multiply by all-company sales.', 'residentialUseGWh': None}},
    }


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('folder', type=Path)
    p.add_argument('--check', action='store_true')
    args = p.parse_args()
    data = build(args.folder)
    sha = emit(OUTPUT, data, args.check)
    manifest = {'schema': data['schema'], 'retrieved': data['retrieved'], 'sha256': sha, 'extract': '../assets/data-center/finland.json', 'inputs': {name: {'sha256': value[0], 'url': value[1]} for name, value in INPUTS.items()}, 'methods': METHODS, 'licenses': {'statfin': {'name': 'CC BY 4.0', 'url': 'https://stat.fi/en/about-us/get-to-know-statistics-finland/legislation/terms-of-use'}, 'energy': 'Source: Energiateollisuus ry. Bounded factual extracts with attribution; no blanket Creative Commons licence asserted. Original workbooks linked, not redistributed.', 'nivos': 'Source: Nivos Vesi ja Lämpö Oy. Limited reported facts with page references; original report linked, not redistributed.'}, 'builder': 'scripts/build_finland_data.py', 'dependency': 'openpyxl==3.1.5'}
    emit(MANIFEST, manifest, args.check)
    print(f'Finnish account: {len(data["prices"]["rows"])} monthly band records; 336 hourly records; heat cells; SHA-256 {sha}')


if __name__ == '__main__':
    main()
